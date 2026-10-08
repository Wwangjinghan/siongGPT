import io
import unittest
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from app.core.domain_types import DocumentBlockType, IngestionErrorCode
from app.parsers.base import ParserLimits
from app.parsers.docx import DocxParser
from app.parsers.errors import IngestionError
from app.parsers.pdf import PdfParser
from app.parsers.registry import ParserRegistry
from app.parsers.text import TextParser
from app.parsers.xlsx import XlsxParser
from app.parsers.zip_safety import inspect_office_zip


def limits(**overrides):
    values = dict(
        max_pdf_pages=10,
        max_extracted_chars=10000,
        max_zip_entries=100,
        max_uncompressed_bytes=1000000,
        max_compression_ratio=100,
        max_xlsx_sheets=5,
        max_xlsx_rows_per_sheet=100,
        max_xlsx_cells_per_row=100,
        max_xlsx_nonempty_cells=1000,
        max_cell_chars=1000,
    )
    values.update(overrides)
    return ParserLimits(**values)


def native_text_pdf(text: str) -> io.BytesIO:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    result = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(f"{number} 0 obj\n".encode())
        result.extend(obj + b"\nendobj\n")
    xref = len(result)
    result.extend(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode())
    result.extend(f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return io.BytesIO(bytes(result))


def docx_fixture() -> io.BytesIO:
    document = Document()
    document.add_heading("Required Documents", level=1)
    document.add_paragraph("Provide PO-001 before approval.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Code"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "GST"
    table.cell(1, 1).text = "Required"
    stream = io.BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream


def xlsx_fixture() -> io.BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Codes"
    sheet.append(["Project Code", "Amount"])
    sheet.append(["PO-001", 100])
    sheet.append(["GST", "=1+1"])
    second = workbook.create_sheet("Other")
    second.append(["DO", "Open"])
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


class TextParserTests(unittest.TestCase):
    def test_utf8_bom_paragraphs_and_line_locators(self):
        blocks = TextParser(limits()).parse(
            io.BytesIO(b"\xef\xbb\xbfFirst\nline\n\nSecond"),
            filename="a.txt",
            mime_type="text/plain",
            source_version_id=uuid4(),
        )
        self.assertEqual([b.content for b in blocks], ["First\nline", "Second"])
        self.assertEqual(blocks[0].locator, {"line_start": 1, "line_end": 2})

    def test_invalid_encoding_and_binary_are_rejected(self):
        parser = TextParser(limits())
        cases = ((b"\xff\xfe", IngestionErrorCode.INVALID_TEXT_ENCODING), (b"a\x00b", IngestionErrorCode.FILE_SIGNATURE_MISMATCH))
        for data, code in cases:
            with self.subTest(code=code), self.assertRaises(IngestionError) as raised:
                parser.parse(io.BytesIO(data), filename="a.txt", mime_type="text/plain", source_version_id=uuid4())
            self.assertEqual(raised.exception.code, code)


class PdfParserTests(unittest.TestCase):
    def test_native_pdf_extracts_page_text(self):
        blocks = PdfParser(limits()).parse(
            native_text_pdf("PO-001 GST"), filename="a.pdf", mime_type="application/pdf", source_version_id=uuid4()
        )
        self.assertIn("PO-001 GST", blocks[0].content)
        self.assertEqual(blocks[0].page, 1)

    def test_blank_pdf_requires_ocr_and_page_limit_is_enforced(self):
        blank = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(blank)
        blank.seek(0)
        with self.assertRaises(IngestionError) as raised:
            PdfParser(limits()).parse(blank, filename="a.pdf", mime_type="application/pdf", source_version_id=uuid4())
        self.assertEqual(raised.exception.code, IngestionErrorCode.OCR_REQUIRED)

        pages = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.add_blank_page(width=100, height=100)
        writer.write(pages)
        pages.seek(0)
        with self.assertRaises(IngestionError) as raised:
            PdfParser(limits(max_pdf_pages=1)).parse(pages, filename="a.pdf", mime_type="application/pdf", source_version_id=uuid4())
        self.assertEqual(raised.exception.code, IngestionErrorCode.PDF_PAGE_LIMIT_EXCEEDED)

    def test_encrypted_and_damaged_pdf_are_rejected(self):
        encrypted = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.encrypt("secret")
        writer.write(encrypted)
        encrypted.seek(0)
        with self.assertRaises(IngestionError) as raised:
            PdfParser(limits()).parse(encrypted, filename="a.pdf", mime_type="application/pdf", source_version_id=uuid4())
        self.assertEqual(raised.exception.code, IngestionErrorCode.ENCRYPTED_PDF)
        with self.assertRaises(IngestionError):
            PdfParser(limits()).parse(io.BytesIO(b"%PDF-broken"), filename="a.pdf", mime_type="application/pdf", source_version_id=uuid4())


class OfficeParserTests(unittest.TestCase):
    def test_docx_preserves_heading_paragraph_and_table(self):
        blocks = DocxParser(limits()).parse(docx_fixture(), filename="a.docx", mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", source_version_id=uuid4())
        self.assertEqual([b.block_type for b in blocks], [DocumentBlockType.HEADING, DocumentBlockType.PARAGRAPH, DocumentBlockType.TABLE])
        self.assertIn("Code | Status", blocks[-1].content)
        self.assertEqual(blocks[1].section, "Required Documents")

    def test_xlsx_preserves_sheets_rows_headers_and_formula_text(self):
        blocks = XlsxParser(limits()).parse(xlsx_fixture(), filename="a.xlsx", mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", source_version_id=uuid4())
        self.assertEqual({b.sheet for b in blocks}, {"Codes", "Other"})
        formula = next(b for b in blocks if "GST" in b.content)
        self.assertIn("=1+1", formula.content)
        self.assertEqual(formula.metadata["table_header"], "Project Code\tAmount")
        self.assertEqual(formula.row_start, 3)

    def test_xlsx_limits_are_enforced(self):
        with self.assertRaises(IngestionError) as raised:
            XlsxParser(limits(max_xlsx_sheets=1)).parse(xlsx_fixture(), filename="a.xlsx", mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", source_version_id=uuid4())
        self.assertEqual(raised.exception.code, IngestionErrorCode.XLSX_LIMIT_EXCEEDED)

        with self.assertRaises(IngestionError) as raised:
            XlsxParser(limits(max_xlsx_cells_per_row=1)).parse(xlsx_fixture(), filename="a.xlsx", mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", source_version_id=uuid4())
        self.assertEqual(raised.exception.code, IngestionErrorCode.XLSX_LIMIT_EXCEEDED)

    def test_zip_bomb_macro_and_corrupt_office_are_rejected(self):
        bomb = io.BytesIO()
        with ZipFile(bomb, "w", ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", "A" * 10000)
        bomb.seek(0)
        with self.assertRaises(IngestionError) as raised:
            inspect_office_zip(bomb, limits(max_compression_ratio=2))
        self.assertEqual(raised.exception.code, IngestionErrorCode.ZIP_BOMB_SUSPECTED)

        macro = io.BytesIO()
        with ZipFile(macro, "w") as archive:
            archive.writestr("word/vbaProject.bin", b"x")
        macro.seek(0)
        with self.assertRaises(IngestionError) as raised:
            inspect_office_zip(macro, limits())
        self.assertEqual(raised.exception.code, IngestionErrorCode.MACRO_CONTENT_REJECTED)

        with self.assertRaises(IngestionError) as raised:
            inspect_office_zip(io.BytesIO(b"PK\x03\x04broken"), limits())
        self.assertEqual(raised.exception.code, IngestionErrorCode.OFFICE_ZIP_INVALID)

    def test_registry_rejects_extension_mime_conflict(self):
        with self.assertRaises(IngestionError) as raised:
            ParserRegistry(limits()).get("fake.pdf", "text/plain")
        self.assertEqual(raised.exception.code, IngestionErrorCode.FILE_SIGNATURE_MISMATCH)
