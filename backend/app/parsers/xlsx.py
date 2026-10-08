from typing import BinaryIO
from uuid import UUID

from openpyxl import load_workbook

from app.core.domain_types import DocumentBlockType, IngestionErrorCode
from app.parsers.base import DocumentBlock, DocumentParser
from app.parsers.errors import IngestionError
from app.parsers.zip_safety import inspect_office_zip


class XlsxParser(DocumentParser):
    name = "openpyxl"

    def parse(self, stream: BinaryIO, *, filename: str, mime_type: str, source_version_id: UUID):
        inspect_office_zip(stream, self.limits)
        workbook = None
        try:
            workbook = load_workbook(
                stream, read_only=True, data_only=False, keep_links=False
            )
            if len(workbook.sheetnames) > self.limits.max_xlsx_sheets:
                raise IngestionError(
                    IngestionErrorCode.XLSX_LIMIT_EXCEEDED,
                    "Workbook exceeds the sheet limit",
                )
            blocks, total_cells, total_chars = [], 0, 0
            for worksheet in workbook.worksheets:
                if worksheet.max_column > self.limits.max_xlsx_cells_per_row:
                    raise IngestionError(
                        IngestionErrorCode.XLSX_LIMIT_EXCEEDED,
                        "Worksheet exceeds the cell-per-row limit",
                    )
                header = None
                for row_no, row in enumerate(worksheet.iter_rows(values_only=True), 1):
                    if row_no > self.limits.max_xlsx_rows_per_sheet:
                        raise IngestionError(
                            IngestionErrorCode.XLSX_LIMIT_EXCEEDED,
                            "Worksheet exceeds the row limit",
                        )
                    values = []
                    for value in row:
                        text = "" if value is None else str(value)
                        if len(text) > self.limits.max_cell_chars:
                            raise IngestionError(
                                IngestionErrorCode.XLSX_LIMIT_EXCEEDED,
                                "Spreadsheet cell exceeds the character limit",
                            )
                        if text:
                            total_cells += 1
                            total_chars += len(text)
                        values.append(text)
                    if total_cells > self.limits.max_xlsx_nonempty_cells:
                        raise IngestionError(
                            IngestionErrorCode.XLSX_LIMIT_EXCEEDED,
                            "Workbook exceeds the non-empty cell limit",
                        )
                    if total_chars > self.limits.max_extracted_chars:
                        raise IngestionError(
                            IngestionErrorCode.EXTRACTED_TEXT_LIMIT_EXCEEDED,
                            "Extracted spreadsheet text exceeds the configured limit",
                        )
                    if not any(values):
                        continue
                    content = "\t".join(values).rstrip("\t")
                    if header is None:
                        header = content
                    blocks.append(
                        DocumentBlock(
                            DocumentBlockType.SPREADSHEET_ROWS,
                            content,
                            len(blocks),
                            sheet=worksheet.title,
                            row_start=row_no,
                            row_end=row_no,
                            locator={
                                "sheet": worksheet.title,
                                "row_start": row_no,
                                "row_end": row_no,
                            },
                            metadata={
                                "parser": self.name,
                                "parser_version": self.version,
                                "table_header": header,
                                "formula_policy": "preserve_formula_text",
                            },
                        )
                    )
            if not blocks:
                raise IngestionError(IngestionErrorCode.EMPTY_DOCUMENT, "XLSX is empty")
            return blocks
        except IngestionError:
            raise
        except Exception as exc:
            raise IngestionError(
                IngestionErrorCode.OFFICE_ZIP_INVALID, "XLSX structure is invalid"
            ) from exc
        finally:
            if workbook is not None:
                workbook.close()
